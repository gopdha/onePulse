"""Report Rendering (HLD Section 2, Step 7; closes FR-5).

FR-5: "The system shall render a finalized report's content into a
locked, per-project visual template." Physical Architecture and LLD
specify this as a deterministic builder with a single hard-coded layout
and no model call — the content it renders was already produced upstream
(Investigation, Synthesis, Deterministic Status Rollup); this module only
lays it out.

Colors are a plain navy-accent identity (no separate brand asset exists
in this repo to pull from) — Arial throughout, per the project's stated
visual direction, not invented from nothing.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from onepulse_common.tower_rollup import TowerRollup

_NAVY = RGBColor(0x1F, 0x38, 0x64)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_DARK_TEXT = RGBColor(0x26, 0x26, 0x26)

_STATUS_COLORS: dict[str, RGBColor] = {
    "Red": RGBColor(0xC0, 0x00, 0x00),
    "Amber": RGBColor(0xED, 0x7D, 0x31),
    "Green": RGBColor(0x54, 0x82, 0x35),
    "Unknown": RGBColor(0x80, 0x80, 0x80),
}

_FONT = "Arial"


@dataclass(frozen=True)
class Finding:
    work_item_id: int
    title: str
    status: str
    evidence: str


def week_of(as_of: _dt.date) -> _dt.date:
    """Monday of the week containing as_of — deterministic, no model call."""
    return as_of - _dt.timedelta(days=as_of.weekday())


def render_status_report(
    *,
    program_name: str,
    as_of: _dt.date,
    overall_status: str,
    executive_summary: str,
    findings: list[Finding],
    output_path: str,
) -> None:
    if overall_status not in _STATUS_COLORS:
        raise ValueError(f"Unrecognized overall_status: {overall_status!r}")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    # -- Header band (navy) with title --------------------------------
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(1.0))
    header.fill.solid()
    header.fill.fore_color.rgb = _NAVY
    header.line.fill.background()
    title_tf = header.text_frame
    title_tf.margin_left = Inches(0.4)
    title_tf.margin_top = Inches(0.1)
    p = title_tf.paragraphs[0]
    p.text = f"{program_name} — Week of {week_of(as_of).isoformat()}"
    p.font.name = _FONT
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = _WHITE

    # -- Overall Status badge ------------------------------------------
    badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.6), Inches(1.25), Inches(2.4), Inches(0.9))
    badge.fill.solid()
    badge.fill.fore_color.rgb = _STATUS_COLORS[overall_status]
    badge.line.fill.background()
    badge_tf = badge.text_frame
    badge_p = badge_tf.paragraphs[0]
    badge_p.text = overall_status.upper()
    badge_p.alignment = PP_ALIGN.CENTER
    badge_p.font.name = _FONT
    badge_p.font.size = Pt(24)
    badge_p.font.bold = True
    badge_p.font.color.rgb = _WHITE

    label_box = slide.shapes.add_textbox(Inches(10.6), Inches(1.0), Inches(2.4), Inches(0.3))
    label_p = label_box.text_frame.paragraphs[0]
    label_p.text = "OVERALL STATUS"
    label_p.alignment = PP_ALIGN.CENTER
    label_p.font.name = _FONT
    label_p.font.size = Pt(11)
    label_p.font.bold = True
    label_p.font.color.rgb = _NAVY

    # -- Executive Summary ----------------------------------------------
    summary_header = slide.shapes.add_textbox(Inches(0.4), Inches(1.25), Inches(9.8), Inches(0.35))
    sh_p = summary_header.text_frame.paragraphs[0]
    sh_p.text = "Executive Summary"
    sh_p.font.name = _FONT
    sh_p.font.size = Pt(16)
    sh_p.font.bold = True
    sh_p.font.color.rgb = _NAVY

    summary_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.65), Inches(9.8), Inches(1.6))
    summary_tf = summary_box.text_frame
    summary_tf.word_wrap = True
    sp = summary_tf.paragraphs[0]
    sp.text = executive_summary
    sp.font.name = _FONT
    sp.font.size = Pt(13)
    sp.font.color.rgb = _DARK_TEXT

    # -- Key Findings -----------------------------------------------------
    kf_header = slide.shapes.add_textbox(Inches(0.4), Inches(3.35), Inches(12.5), Inches(0.35))
    kf_p = kf_header.text_frame.paragraphs[0]
    kf_p.text = "Key Findings"
    kf_p.font.name = _FONT
    kf_p.font.size = Pt(16)
    kf_p.font.bold = True
    kf_p.font.color.rgb = _NAVY

    findings_box = slide.shapes.add_textbox(Inches(0.4), Inches(3.75), Inches(12.5), Inches(3.4))
    findings_tf = findings_box.text_frame
    findings_tf.word_wrap = True

    for i, finding in enumerate(findings):
        title_p = findings_tf.paragraphs[0] if i == 0 else findings_tf.add_paragraph()
        title_p.text = f"#{finding.work_item_id}  {finding.title}  —  {finding.status}"
        title_p.font.name = _FONT
        title_p.font.size = Pt(14)
        title_p.font.bold = True
        title_p.font.color.rgb = _DARK_TEXT

        evidence_p = findings_tf.add_paragraph()
        evidence_p.text = f"Evidence: {finding.evidence}"
        evidence_p.level = 1
        evidence_p.font.name = _FONT
        evidence_p.font.size = Pt(12)
        evidence_p.font.italic = True
        evidence_p.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    prs.save(output_path)


_TOWER_CARD_BG = RGBColor(0xF7, 0xF8, 0xFA)
_MUTED_TEXT = RGBColor(0x6B, 0x70, 0x78)


def render_tower_report(
    *,
    program_name: str,
    as_of: _dt.date,
    program_health_status: str,
    towers: list[TowerRollup],
    untracked_initiatives: list[dict],
    output_path: str,
) -> None:
    """Executive Tower View (Task 36) — real, deterministic layout built
    from `tower_rollup.build_tower_rollups`'s real per-tower/per-feature
    numbers. Only ever called when `towers` is non-empty; the caller
    (`pipeline.run_pipeline_cycle`) falls back to `render_status_report`
    above, unchanged, when a project has no real Epic parents above its
    Committed Features.

    Real, deliberate layout approximations from the reference mockup,
    stated plainly rather than silently settled for: feature dots are
    plain colored bullet glyphs, not separate drawn shapes (python-pptx
    text runs are the simpler, equally legible real mechanism here); and
    real ADO titles keep their raw sequence-number prefixes (e.g. "07
    Build OTEL Collector - AWS") rather than being stripped to match the
    mockup's hand-cleaned text — a regex-based cleanup would be a
    fragile guess at one organization's own naming convention, not a
    real, general rule, so the honest choice is to render the real title
    as-is.
    """
    if program_health_status not in _STATUS_COLORS:
        raise ValueError(f"Unrecognized program_health_status: {program_health_status!r}")
    if not towers:
        raise ValueError("render_tower_report requires at least one real TowerRollup — use render_status_report for the flat fallback instead.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # -- Header band (navy) ---------------------------------------------
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(1.1))
    header.fill.solid()
    header.fill.fore_color.rgb = _NAVY
    header.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.13), Inches(8.0), Inches(0.5))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = program_name
    title_p.font.name = _FONT
    title_p.font.size = Pt(24)
    title_p.font.bold = True
    title_p.font.color.rgb = _WHITE

    subtitle_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.62), Inches(8.0), Inches(0.35))
    subtitle_p = subtitle_box.text_frame.paragraphs[0]
    subtitle_p.text = f"Week of {week_of(as_of).isoformat()}"
    subtitle_p.font.name = _FONT
    subtitle_p.font.size = Pt(13)
    subtitle_p.font.color.rgb = RGBColor(0xB9, 0xC4, 0xD9)

    health_label = slide.shapes.add_textbox(Inches(10.4), Inches(0.12), Inches(2.5), Inches(0.25))
    health_label_p = health_label.text_frame.paragraphs[0]
    health_label_p.text = "PROGRAM HEALTH"
    health_label_p.alignment = PP_ALIGN.RIGHT
    health_label_p.font.name = _FONT
    health_label_p.font.size = Pt(9)
    health_label_p.font.bold = True
    health_label_p.font.color.rgb = RGBColor(0xB9, 0xC4, 0xD9)

    health_badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.9), Inches(0.4), Inches(2.0), Inches(0.5))
    health_badge.fill.solid()
    health_badge.fill.fore_color.rgb = _STATUS_COLORS[program_health_status]
    health_badge.line.fill.background()
    health_badge_p = health_badge.text_frame.paragraphs[0]
    health_badge_p.text = program_health_status.upper()
    health_badge_p.alignment = PP_ALIGN.CENTER
    health_badge_p.font.name = _FONT
    health_badge_p.font.size = Pt(18)
    health_badge_p.font.bold = True
    health_badge_p.font.color.rgb = _WHITE

    # -- Tower cards, up to 3 side by side --------------------------------
    card_top = Inches(1.35)
    card_height = Inches(2.75)
    card_width = Inches(4.0)
    card_gap = Inches(0.25)
    left_margin = Inches(0.4)

    for i, tower in enumerate(towers[:3]):
        card_left = left_margin + i * (card_width + card_gap)

        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, card_left, card_top, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = _TOWER_CARD_BG
        card.line.fill.background()

        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, card_left, card_top, card_width, Inches(0.06))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = _STATUS_COLORS[tower.health]
        top_bar.line.fill.background()

        title_tb = slide.shapes.add_textbox(card_left + Inches(0.2), card_top + Inches(0.15), card_width - Inches(0.4), Inches(0.6))
        title_tf = title_tb.text_frame
        title_tf.word_wrap = True
        t_p = title_tf.paragraphs[0]
        t_p.text = tower.title
        t_p.font.name = _FONT
        t_p.font.size = Pt(15)
        t_p.font.bold = True
        t_p.font.color.rgb = _NAVY

        pct_tb = slide.shapes.add_textbox(card_left + Inches(0.2), card_top + Inches(0.75), Inches(1.5), Inches(0.75))
        pct_p = pct_tb.text_frame.paragraphs[0]
        pct_p.text = f"{tower.percent_complete}%"
        pct_p.font.name = _FONT
        pct_p.font.size = Pt(32)
        pct_p.font.bold = True
        pct_p.font.color.rgb = _STATUS_COLORS[tower.health]

        pct_caption_tb = slide.shapes.add_textbox(card_left + Inches(0.2), card_top + Inches(1.42), Inches(1.5), Inches(0.25))
        pct_caption_p = pct_caption_tb.text_frame.paragraphs[0]
        pct_caption_p.text = "complete"
        pct_caption_p.font.name = _FONT
        pct_caption_p.font.size = Pt(9)
        pct_caption_p.font.color.rgb = _MUTED_TEXT

        stats_tb = slide.shapes.add_textbox(card_left + Inches(1.8), card_top + Inches(0.8), card_width - Inches(2.0), Inches(0.65))
        stats_tf = stats_tb.text_frame
        stats_tf.word_wrap = True
        delivered_p = stats_tf.paragraphs[0]
        delivered_p.text = f"{tower.delivered_count} of {tower.total_count} items delivered"
        delivered_p.font.name = _FONT
        delivered_p.font.size = Pt(10)
        delivered_p.font.color.rgb = _MUTED_TEXT
        if tower.flagged_count > 0:
            flagged_p = stats_tf.add_paragraph()
            noun = "needs" if tower.flagged_count == 1 else "need"
            flagged_p.text = f"{tower.flagged_count} {noun} an owner"
            flagged_p.font.name = _FONT
            flagged_p.font.size = Pt(10)
            flagged_p.font.bold = True
            flagged_p.font.color.rgb = _STATUS_COLORS["Amber"]

        features_tb = slide.shapes.add_textbox(card_left + Inches(0.2), card_top + Inches(1.75), card_width - Inches(0.4), card_height - Inches(1.9))
        features_tf = features_tb.text_frame
        features_tf.word_wrap = True
        for j, feature in enumerate(tower.features):
            f_p = features_tf.paragraphs[0] if j == 0 else features_tf.add_paragraph()
            # Two real, separate runs — the leading dot glyph colored by
            # this feature's own real health, the title itself kept plain
            # dark text regardless of health (matching the reference:
            # only the dot signals status).
            dot_run = f_p.add_run()
            dot_run.text = "●  "
            dot_run.font.color.rgb = _STATUS_COLORS[feature.health]
            title_run = f_p.add_run()
            title_run.text = feature.title
            title_run.font.color.rgb = _DARK_TEXT
            for run in (dot_run, title_run):
                run.font.name = _FONT
                run.font.size = Pt(10.5)

    # -- Two columns: This Week's Initiatives / Needs Your Decision ------
    columns_top = Inches(4.35)
    col_width = Inches(6.1)

    init_header = slide.shapes.add_textbox(Inches(0.4), columns_top, col_width, Inches(0.35))
    init_header_p = init_header.text_frame.paragraphs[0]
    init_header_p.text = "This Week's Initiatives"
    init_header_p.font.name = _FONT
    init_header_p.font.size = Pt(15)
    init_header_p.font.bold = True
    init_header_p.font.color.rgb = _NAVY

    init_sub = slide.shapes.add_textbox(Inches(0.4), columns_top + Inches(0.35), col_width, Inches(0.3))
    init_sub_p = init_sub.text_frame.paragraphs[0]
    init_sub_p.text = "Work happening alongside the towers above — not yet reflected in delivery tracking."
    init_sub_p.font.name = _FONT
    init_sub_p.font.size = Pt(9)
    init_sub_p.font.italic = True
    init_sub_p.font.color.rgb = _MUTED_TEXT

    init_box = slide.shapes.add_textbox(Inches(0.4), columns_top + Inches(0.75), col_width, Inches(2.2))
    init_tf = init_box.text_frame
    init_tf.word_wrap = True
    if untracked_initiatives:
        for i, initiative in enumerate(untracked_initiatives):
            p = init_tf.paragraphs[0] if i == 0 else init_tf.add_paragraph()
            # Real, near-verbatim rendering (Task 36 decision 2): the
            # description and evidence come straight from Status
            # Analysis's own real output, no rephrasing step.
            r1 = p.add_run()
            r1.text = f"{initiative['description']} — "
            r1.font.bold = True
            r2 = p.add_run()
            r2.text = initiative.get("evidence", "")
            for run in p.runs:
                run.font.name = _FONT
                run.font.size = Pt(11)
                run.font.color.rgb = _DARK_TEXT if run is r1 else _MUTED_TEXT
    else:
        p = init_tf.paragraphs[0]
        p.text = "None this week."
        p.font.name = _FONT
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = _MUTED_TEXT

    decision_lines = [f.decision_line for t in towers for f in t.features if f.decision_line]

    dec_header = slide.shapes.add_textbox(Inches(6.85), columns_top, col_width, Inches(0.35))
    dec_header_p = dec_header.text_frame.paragraphs[0]
    dec_header_p.text = "Needs Your Decision"
    dec_header_p.font.name = _FONT
    dec_header_p.font.size = Pt(15)
    dec_header_p.font.bold = True
    dec_header_p.font.color.rgb = _NAVY

    dec_sub = slide.shapes.add_textbox(Inches(6.85), columns_top + Inches(0.35), col_width, Inches(0.3))
    dec_sub_p = dec_sub.text_frame.paragraphs[0]
    dec_sub_p.text = "Delivered work waiting on an owner to ship."
    dec_sub_p.font.name = _FONT
    dec_sub_p.font.size = Pt(9)
    dec_sub_p.font.italic = True
    dec_sub_p.font.color.rgb = _MUTED_TEXT

    dec_box = slide.shapes.add_textbox(Inches(6.85), columns_top + Inches(0.75), col_width, Inches(2.2))
    dec_tf = dec_box.text_frame
    dec_tf.word_wrap = True
    if decision_lines:
        for i, line in enumerate(decision_lines):
            p = dec_tf.paragraphs[0] if i == 0 else dec_tf.add_paragraph()
            p.text = f"•  {line}"
            p.font.name = _FONT
            p.font.size = Pt(11)
            p.font.color.rgb = _DARK_TEXT
    else:
        p = dec_tf.paragraphs[0]
        p.text = "None this week — every delivered item has an owner."
        p.font.name = _FONT
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = _MUTED_TEXT

    prs.save(output_path)
