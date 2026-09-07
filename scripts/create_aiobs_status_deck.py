"""One-off script (Task 27 stress test): writes a minimal, real status
deck for the "Agentic AI Observability Platform" project so
run_pipeline.py's Status Update Analysis stage has a real .pptx to read.
Not a curated demo like Leave Tracker's — this stress test's real focus
is Investigation's behavior at 670 real work items, not Status Analysis
quality, so this deck is intentionally minimal: it references two real
epics by their actual ID/title (epic 218 "Observability for Agentic
Frameworks", epic 530 "AI Agent Operations Center") plus one
deliberately untracked initiative, the same honest pattern
scripts/seed_leave_tracker_project.py already established.
"""

from pptx import Presentation
from pptx.util import Inches, Pt

output_path = "aiobs_status_deck.pptx"

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.6))
title_box.text_frame.paragraphs[0].text = "Team Lead Status Update — Agentic AI Observability Platform — Week of Sept 6"
title_box.text_frame.paragraphs[0].font.size = Pt(22)
title_box.text_frame.paragraphs[0].font.bold = True

body_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5))
tf = body_box.text_frame
tf.word_wrap = True

bullets = [
    "Observability for Agentic Frameworks (Epic #218): tracing and eval "
    "instrumentation work continuing across several features this sprint.",
    "AI Agent Operations Center (Epic #530): operations tooling underway; "
    "several features still in early design.",
    "A separate vendor security review for our telemetry backend is in "
    "progress with procurement — not yet tracked as an Azure DevOps work "
    "item.",
]
for i, bullet in enumerate(bullets):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = f"• {bullet}"
    p.font.size = Pt(15)

prs.save(output_path)
print(f"Wrote {output_path}")
