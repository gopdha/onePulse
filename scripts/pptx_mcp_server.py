"""Custom PPTX-parsing MCP server for Status Update Analysis (FR-2).

Physical Architecture Section 2 names this exact component: "Status
Update Analysis | AKS service; Claude via Foundry; custom PPTX-parsing
MCP server." This is that server — a real MCP server (not a spike
against a third-party one), run as a local stdio subprocess exactly like
the official Azure DevOps MCP server, so Status Analysis can be bridged
into an AgentsClient agent the same proven way as Investigation.

Uses the mcp SDK's 2.x server API: FastMCP was renamed to MCPServer in
mcp>=2.0 (discovered live — the 1.x `from mcp.server.fastmcp import
FastMCP` import raises ModuleNotFoundError on the installed mcp==2.1.1
with a migration pointer to `mcp.server.mcpserver.MCPServer`).

Two tools only — least-privilege (NFR-3): this server can read slide
text, nothing else. No write tools, no filesystem access beyond a single
named .pptx path per call.

Run standalone for a manual check: python scripts/pptx_mcp_server.py
Run as a subprocess: this is what scripts/status_analysis_spike.py does.
"""

from __future__ import annotations

from pptx import Presentation

from mcp.server.mcpserver import MCPServer

server = MCPServer(name="pptx-parser")


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


@server.tool()
def list_slides(pptx_path: str) -> list[dict]:
    """List every slide in the given .pptx file with its index and a short text preview."""
    prs = Presentation(pptx_path)
    return [
        {"slide_index": i, "preview": _slide_text(slide)[:120]}
        for i, slide in enumerate(prs.slides)
    ]


@server.tool()
def get_slide_text(pptx_path: str, slide_index: int) -> str:
    """Return the full text content of one slide (all text frames, in shape order)."""
    prs = Presentation(pptx_path)
    if slide_index < 0 or slide_index >= len(prs.slides):
        raise ValueError(f"slide_index {slide_index} out of range (deck has {len(prs.slides)} slides)")
    return _slide_text(prs.slides[slide_index])


if __name__ == "__main__":
    import asyncio

    asyncio.run(server.run_stdio_async())
