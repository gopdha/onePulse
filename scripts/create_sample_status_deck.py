"""Creates a small, real sample team-lead status deck for testing Status
Update Analysis (FR-2) against real PPTX content. Not a stub: it's a real
.pptx file, parsed by a real MCP server, reasoned about by a real agent.

Deliberately includes:
  - Text overlapping real Azure DevOps items 8/9/10 in the singleSlide
    project ("Coding", "Analysis") — should be flagged as possible
    connections to tracked work.
  - One initiative with no ADO counterpart ("vendor contract renewal")
    — should be flagged as untracked, per FR-2.

Run once: python scripts/create_sample_status_deck.py
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt

OUTPUT_PATH = "sample_status_deck.pptx"


def main() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.6))
    title_box.text_frame.paragraphs[0].text = "Team Lead Status Update — Week of Sept 1"
    title_box.text_frame.paragraphs[0].font.size = Pt(24)
    title_box.text_frame.paragraphs[0].font.bold = True

    body_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(4.5))
    tf = body_box.text_frame
    tf.word_wrap = True

    bullets = [
        "Coding: implementation underway, on pace for end of sprint.",
        "Analysis: wrapping up this week, findings being written up.",
        "Vendor contract renewal: blocked on external counsel review, "
        "no ETA yet from Legal. Not currently tracked as an Azure DevOps work item.",
    ]
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.size = Pt(16)

    prs.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
